"""Тесты извлечения текста из файлов реквизитов (requisites_extract)."""
from __future__ import annotations

from pathlib import Path

import pytest

from src.contracts.requisites_extract import (
    NoTextLayerError,
    extract_text,
)

_INN = "7707083893"


def _make_docx_with_table(path: Path, rows: list[tuple[str, str]]) -> None:
    import docx

    doc = docx.Document()
    table = doc.add_table(rows=len(rows), cols=2)
    for i, (left, right) in enumerate(rows):
        table.cell(i, 0).text = left
        table.cell(i, 1).text = right
    doc.save(str(path))


def _make_pdf(path: Path, text: str | None) -> None:
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    if text:
        page.insert_text((72, 72), text)
    doc.save(str(path))
    doc.close()


class TestExtractDocx:
    def test_table_cells_extracted(self, tmp_path: Path):
        """ИНН из ячейки таблицы попадает в извлечённый текст."""
        path = tmp_path / "card.docx"
        _make_docx_with_table(path, [
            ("Наименование", 'ООО "Ромашка"'),
            ("ИНН", _INN),
        ])
        text = extract_text(path)
        assert _INN in text
        assert "Ромашка" in text

    def test_short_valid_docx_not_flagged_as_scan(self, tmp_path: Path):
        """Короткий валидный docx (< 50 символов) — текст, не NoTextLayerError."""
        path = tmp_path / "short.docx"
        _make_docx_with_table(path, [("ИНН", _INN)])
        text = extract_text(path)
        assert _INN in text
        assert len("".join(text.split())) < 50


class TestExtractPdf:
    def test_text_layer_extracted(self, tmp_path: Path):
        path = tmp_path / "card.pdf"
        _make_pdf(path, f'ООО "Ромашка", ИНН {_INN}, ОГРН 1027700132195, г. Москва')
        text = extract_text(path)
        assert _INN in text

    def test_scan_without_text_layer_raises(self, tmp_path: Path):
        """PDF без текста (пустая страница) → NoTextLayerError."""
        path = tmp_path / "scan.pdf"
        _make_pdf(path, None)
        with pytest.raises(NoTextLayerError):
            extract_text(path)


class TestDispatcher:
    def test_unsupported_extension_raises(self, tmp_path: Path):
        path = tmp_path / "card.xlsx"
        path.write_bytes(b"")
        with pytest.raises(ValueError):
            extract_text(path)
